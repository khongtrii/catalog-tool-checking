from __future__ import annotations

from io import BytesIO
import unittest
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook, Workbook
import polars as pl

import catalog_checker
from catalog_checker import (
    WorkbookValidationError,
    _unique_normalized_column_mapping,
    results_to_excel,
    validate_xlsx,
)


def workbook_bytes() -> bytes:
    workbook = Workbook()
    workbook.active["A1"] = "VALUE"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class WorkbookSecurityTests(unittest.TestCase):
    def test_normalized_headers_are_unique(self) -> None:
        mapping = _unique_normalized_column_mapping(["START DATE", "START/DATE", ""])
        self.assertEqual(
            list(mapping.values()),
            ["START DATE", "START DATE 2", "COLUMN"],
        )

    def test_valid_workbook_is_accepted(self) -> None:
        validate_xlsx(workbook_bytes())

    def test_non_xlsx_is_rejected(self) -> None:
        with self.assertRaises(WorkbookValidationError):
            validate_xlsx(b"not an xlsx file")

    def test_oversized_upload_is_rejected(self) -> None:
        with patch.object(catalog_checker, "MAX_XLSX_BYTES", 4):
            with self.assertRaisesRegex(WorkbookValidationError, "20 MB"):
                validate_xlsx(workbook_bytes())

    def test_large_uncompressed_container_is_rejected(self) -> None:
        output = BytesIO()
        with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "content")
            archive.writestr("xl/workbook.xml", "workbook")
            archive.writestr("xl/worksheets/sheet1.xml", "x" * 100)
        with patch.object(catalog_checker, "MAX_XLSX_UNCOMPRESSED_BYTES", 32):
            with self.assertRaisesRegex(WorkbookValidationError, "giải nén"):
                validate_xlsx(output.getvalue())

    def test_formula_like_text_is_escaped_in_output(self) -> None:
        frame = pl.DataFrame(
            {
                "VALUE": ["=1+1", "+SUM(A1:A2)", "-2+3", "@payload", "safe"],
            }
        )
        output = results_to_excel([("Result", frame)])
        workbook = load_workbook(BytesIO(output), data_only=False)
        values = [workbook.active.cell(row=row, column=1).value for row in range(2, 7)]
        self.assertEqual(values, ["'=1+1", "'+SUM(A1:A2)", "'-2+3", "'@payload", "safe"])
        self.assertTrue(all(workbook.active.cell(row=row, column=1).data_type != "f" for row in range(2, 7)))


if __name__ == "__main__":
    unittest.main()
