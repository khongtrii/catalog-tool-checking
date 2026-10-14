from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re
import unicodedata
from zipfile import ZIP_DEFLATED, ZipFile

import polars as pl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


REQUIRED_CATALOG_COLUMNS = [
    "GROUPED DIVISION",
    "DIVISION",
    "DEPARTMENT",
    "GOLD CODE",
    "TILLCODE",
    "SV",
    "PRODUCT NAME VIE",
    "UNIT",
    "START DATE",
    "END DATE",
    "SYSTEM GOLD VLL",
    "REGION NORTH SOUTH CENTER ALL",
    "PAGE",
    "PAGE DIGITAL",
    "THEMATIC",
    "POSITION",
    "NORMAL SELLING PRICE VND",
    "PROMOTIONAL SELLING PRICE VND",
    "MULTIBUYCOMBO SELLING PRICE VND",
    "DISCOUNT TO CUSTOMER VND",
    "% DISCOUNT TO CUSTOMER",
    "PROMO TYPE",
    "PROMO DETAILS",
    "PROMOTION GIFT",
    "GIFT VALUE",
    "NOTE",
    "THE1 POINTS",
    "THE1 EXCLUSIVE PRICE VND",
    "NOTE_1",
]

HEADER_COLUMNS = {
    "GROUPED DIVISION",
    "DIVISION",
    "DEPARTMENT",
    "GOLD CODE",
    "TILLCODE",
    "SV",
}

QUERY_COLUMNS = {
    "DIVISION",
    "CH",
    "GOLD_CODE",
    "SV",
    "GIA_NORMAL",
    "PRICELIST_NORMAL",
    "START_DATE",
    "END_DATE",
    "GIA_KM",
    "PRICELIST_KM",
    "START_DATE_KM",
    "END_DATE_KM",
}

BUSINESS_GROUPS = ("NON FOOD", "FMCG", "FRESH", "OTHERS")


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    normalized = unicodedata.normalize("NFKC", str(value))
    normalized = re.sub(r"[/\r\n\t\f\v]+", " ", normalized)
    normalized = re.sub(r"[^\w\s%]", "", normalized)
    normalized = re.sub(r"\s+", " ", normalized)
    return normalized.strip().upper()


def parse_target_date(value: date | datetime | str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for date_format in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, date_format).date()
        except ValueError:
            continue
    raise ValueError("Ngày kiểm tra phải có định dạng DD/MM/YYYY.")


def workbook_sheet_info(file_bytes: bytes) -> list[dict[str, object]]:
    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    result = []
    try:
        for worksheet in workbook.worksheets:
            header_row = _find_header_row_in_worksheet(worksheet)
            result.append(
                {
                    "name": worksheet.title,
                    "header_row": header_row,
                    "processable": header_row is not None,
                }
            )
    finally:
        workbook.close()
    return result


def _find_header_row_in_worksheet(worksheet, max_rows: int = 20) -> int | None:
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(max_rows, worksheet.max_row), values_only=True)
    ):
        values = {normalize_header(value) for value in row if value is not None}
        if HEADER_COLUMNS.issubset(values):
            return row_index
    return None


def _catalog_column_mapping(dataframe: pl.DataFrame) -> pl.DataFrame:
    mapping: dict[str, str] = {}
    used: set[str] = set()

    for required in REQUIRED_CATALOG_COLUMNS:
        if required in dataframe.columns:
            used.add(required)
            continue
        matches = [
            column
            for column in dataframe.columns
            if column.startswith(required + " ") and column not in used
        ]
        if len(matches) == 1:
            mapping[matches[0]] = required
            used.add(matches[0])

    dataframe = dataframe.rename(mapping)
    return dataframe.select(
        [
            pl.col(column)
            if column in dataframe.columns
            else pl.lit(None, dtype=pl.String).alias(column)
            for column in REQUIRED_CATALOG_COLUMNS
        ]
    )


def _identifier(column: str) -> pl.Expr:
    return (
        pl.col(column)
        .cast(pl.String)
        .str.strip_chars()
        .str.replace(r"\.0$", "")
    )


def _price(column: str) -> pl.Expr:
    return (
        pl.col(column)
        .cast(pl.String)
        .str.replace_all(r"[,\s]", "")
        .cast(pl.Float64, strict=False)
        .round(0)
        .cast(pl.Int64)
    )


def _date_column(column: str) -> pl.Expr:
    text = pl.col(column).cast(pl.String).str.strip_chars()
    return pl.coalesce(
        text.str.strptime(pl.Date, format="%d/%m/%y", strict=False),
        text.str.strptime(pl.Date, format="%d/%m/%Y", strict=False),
        text.str.strptime(pl.Date, format="%Y-%m-%d", strict=False),
        text.str.strptime(pl.Date, format="%Y-%m-%d %H:%M:%S", strict=False),
    )


def read_catalog_sheet(
    file_bytes: bytes,
    sheet_name: str,
    suggest_site: str,
    header_row: int | None = None,
) -> pl.DataFrame:
    if header_row is None:
        sheet_info = {item["name"]: item for item in workbook_sheet_info(file_bytes)}
        header_row = sheet_info.get(sheet_name, {}).get("header_row")
    if header_row is None:
        raise ValueError(f"Không tìm thấy header catalogue trong sheet '{sheet_name}'.")

    dataframe = pl.read_excel(
        BytesIO(file_bytes),
        sheet_name=sheet_name,
        infer_schema_length=0,
        read_options={"header_row": header_row},
    )
    dataframe = dataframe.rename({column: normalize_header(column) for column in dataframe.columns})
    dataframe = _catalog_column_mapping(dataframe)
    site = str(suggest_site).strip()
    if not site:
        raise ValueError(f"REGION '{sheet_name}' chưa có SUGGEST SITE.")

    dataframe = dataframe.with_columns(
        _identifier("GOLD CODE"),
        _identifier("SV"),
        _price("NORMAL SELLING PRICE VND"),
        _price("PROMOTIONAL SELLING PRICE VND"),
    )
    dataframe = dataframe.filter(
        pl.col("GOLD CODE").is_not_null()
        & (pl.col("GOLD CODE") != "")
        & pl.col("SV").is_not_null()
        & (pl.col("SV") != "")
        & (
            pl.col("NORMAL SELLING PRICE VND").is_not_null()
            | pl.col("PROMOTIONAL SELLING PRICE VND").is_not_null()
        )
    )
    return dataframe.with_columns(
        pl.lit(sheet_name).alias("REGION"),
        pl.lit(site).alias("SUGGEST SITE"),
        pl.coalesce(
            "PROMOTIONAL SELLING PRICE VND",
            "NORMAL SELLING PRICE VND",
        ).alias("FINAL PRICE"),
        pl.concat_str(
            [pl.col("GOLD CODE"), pl.col("SV"), pl.lit(site)],
            separator="-",
        ).alias("KEY"),
    )


def read_effective_query(file_bytes: bytes, target_date: date | datetime | str) -> pl.DataFrame:
    target_date = parse_target_date(target_date)
    query = pl.read_excel(BytesIO(file_bytes), infer_schema_length=0)
    missing = sorted(QUERY_COLUMNS.difference(query.columns))
    if missing:
        raise ValueError("File query thiếu cột: " + ", ".join(missing))

    query = query.filter(pl.col("DIVISION") != "DIVISION").with_columns(
        _identifier("GOLD_CODE"),
        _identifier("SV"),
        _identifier("CH"),
        _identifier("PRICELIST_NORMAL"),
        _identifier("PRICELIST_KM"),
        _price("GIA_NORMAL").alias("_NORMAL_PRICE"),
        _price("GIA_KM").alias("_PROMO_PRICE"),
        _date_column("START_DATE").alias("_NORMAL_START_DATE"),
        _date_column("END_DATE").alias("_NORMAL_END_DATE"),
        _date_column("START_DATE_KM").alias("_PROMO_START_DATE"),
        _date_column("END_DATE_KM").alias("_PROMO_END_DATE"),
    )

    normal_is_valid = (
        pl.col("_NORMAL_PRICE").is_not_null()
        & (pl.col("_NORMAL_START_DATE") <= pl.lit(target_date))
        & (pl.col("_NORMAL_END_DATE") >= pl.lit(target_date))
    )
    promo_is_valid = (
        pl.col("_PROMO_PRICE").is_not_null()
        & (pl.col("_PROMO_START_DATE") <= pl.lit(target_date))
        & (pl.col("_PROMO_END_DATE") >= pl.lit(target_date))
    )

    query = query.with_columns(
        pl.when(promo_is_valid)
        .then(pl.col("PRICELIST_KM"))
        .when(normal_is_valid)
        .then(pl.col("PRICELIST_NORMAL"))
        .alias("QUERY PRICE LIST"),
        pl.when(promo_is_valid)
        .then(pl.col("_PROMO_PRICE"))
        .when(normal_is_valid)
        .then(pl.col("_NORMAL_PRICE"))
        .alias("QUERY FINAL PRICE"),
        pl.when(promo_is_valid)
        .then(pl.col("_PROMO_START_DATE"))
        .when(normal_is_valid)
        .then(pl.col("_NORMAL_START_DATE"))
        .alias("QUERY VALID START DATE"),
        pl.when(promo_is_valid)
        .then(pl.col("_PROMO_END_DATE"))
        .when(normal_is_valid)
        .then(pl.col("_NORMAL_END_DATE"))
        .alias("QUERY VALID END DATE"),
        pl.when(promo_is_valid)
        .then(pl.lit("KM"))
        .when(normal_is_valid)
        .then(pl.lit("NORMAL"))
        .alias("QUERY PRICE TYPE"),
        pl.concat_str(["GOLD_CODE", "SV", "CH"], separator="-").alias("KEY"),
    ).filter(pl.col("QUERY FINAL PRICE").is_not_null())

    conflicts = (
        query.group_by("KEY")
        .agg(pl.col("QUERY FINAL PRICE").n_unique().alias("PRICE COUNT"))
        .filter(pl.col("PRICE COUNT") > 1)
    )
    if conflicts.height:
        examples = ", ".join(conflicts.head(5)["KEY"].to_list())
        raise ValueError(f"Query có nhiều giá hiệu lực khác nhau cho cùng KEY: {examples}")

    return (
        query.select(
            "KEY",
            "QUERY PRICE LIST",
            "QUERY FINAL PRICE",
            "QUERY PRICE TYPE",
            "QUERY VALID START DATE",
            "QUERY VALID END DATE",
        )
        .unique(subset="KEY", keep="first", maintain_order=True)
    )


def compare_with_query(catalog: pl.DataFrame, query: pl.DataFrame) -> pl.DataFrame:
    return catalog.join(query, on="KEY", how="inner").with_columns(
        pl.when(pl.col("FINAL PRICE") == pl.col("QUERY FINAL PRICE"))
        .then(pl.lit("đúng giá"))
        .otherwise(pl.lit("lệch giá"))
        .alias("MDT NOTE")
    )


def _division_code(column: str) -> pl.Expr:
    value = pl.col(column).cast(pl.String).str.strip_chars()
    two_digit_code = value.str.extract(r"^(0[1-6])", 1)
    one_digit_code = value.str.extract(r"^([1-6])", 1)
    return pl.coalesce(
        two_digit_code,
        pl.concat_str(pl.lit("0"), one_digit_code),
    )


def split_by_business_group(
    dataframe: pl.DataFrame,
    primary_column: str,
    fallback_column: str | None = None,
) -> dict[str, pl.DataFrame]:
    if primary_column not in dataframe.columns:
        if fallback_column and fallback_column in dataframe.columns:
            primary_column = fallback_column
            fallback_column = None
        else:
            return {"OTHERS": dataframe}

    code = _division_code(primary_column)
    if fallback_column and fallback_column in dataframe.columns:
        fallback_code = _division_code(fallback_column)
        code = pl.when(code.is_in(["01", "02", "03", "04", "05", "06"])).then(
            code
        ).otherwise(fallback_code)

    classified = dataframe.with_columns(
        pl.when(code.is_in(["01", "02", "03"]))
        .then(pl.lit("NON FOOD"))
        .when(code.is_in(["04", "06"]))
        .then(pl.lit("FMCG"))
        .when(code == "05")
        .then(pl.lit("FRESH"))
        .otherwise(pl.lit("OTHERS"))
        .alias("_BUSINESS_GROUP")
    )
    return {
        group: subset.drop("_BUSINESS_GROUP")
        for group in BUSINESS_GROUPS
        if (
            subset := classified.filter(pl.col("_BUSINESS_GROUP") == group)
        ).height
    }


def derive_output_sheet_name(filename: str) -> str:
    stem = filename.rsplit(".", 1)[0].strip()
    match = re.match(r"(?i)^(BF\d+|C\d+|FLYER)\b", stem)
    return match.group(1).upper() if match else stem[:31]


def _safe_sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", name).strip(" '") or "Result"
    base = base[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        marker = f"_{suffix}"
        candidate = base[: 31 - len(marker)] + marker
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def results_to_excel(results: list[tuple[str, pl.DataFrame]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()
    body_font = Font(name="Arial", size=10, color="1F1F1F")
    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    for requested_name, dataframe in results:
        worksheet = workbook.create_sheet(_safe_sheet_name(requested_name, used_names))
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        worksheet.append(dataframe.columns)
        for cell in worksheet[1]:
            fill_color, font_color = _header_colors(str(cell.value))
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.font = Font(
                name="Arial",
                size=10,
                bold=True,
                color=font_color,
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = thin_border
        worksheet.row_dimensions[1].height = 72

        widths = [len(str(column)) for column in dataframe.columns]
        for row in dataframe.iter_rows():
            worksheet.append(row)
            for index, value in enumerate(row):
                if value is not None:
                    widths[index] = min(max(widths[index], len(str(value))), 40)

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="center")
        for index, width in enumerate(widths, start=1):
            column_name = str(worksheet.cell(1, index).value)
            if column_name.startswith("QUERY "):
                target_width = 16
            elif column_name in {
                "PRODUCT NAME VIE",
                "PROMO DETAILS",
                "NOTE",
                "NOTE_1",
            }:
                target_width = min(max(18, width + 2), 32)
            else:
                target_width = min(max(11, width + 2), 18)
            worksheet.column_dimensions[
                worksheet.cell(1, index).column_letter
            ].width = target_width

        worksheet.auto_filter.ref = worksheet.dimensions
        for column_name in ("FINAL PRICE", "QUERY FINAL PRICE"):
            if column_name in dataframe.columns:
                column_index = dataframe.columns.index(column_name) + 1
                for cell in worksheet.iter_cols(
                    min_col=column_index,
                    max_col=column_index,
                    min_row=2,
                    max_row=worksheet.max_row,
                ):
                    for item in cell:
                        item.number_format = "#,##0"
        for column_name in ("QUERY VALID START DATE", "QUERY VALID END DATE"):
            if column_name in dataframe.columns:
                column_index = dataframe.columns.index(column_name) + 1
                for cells in worksheet.iter_cols(
                    min_col=column_index,
                    max_col=column_index,
                    min_row=2,
                    max_row=worksheet.max_row,
                ):
                    for item in cells:
                        item.number_format = "dd/mm/yyyy"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _header_colors(column_name: str) -> tuple[str, str]:
    if column_name.startswith("QUERY "):
        return "48ACD3", "FFFFFF"
    if column_name in {"THE1 POINTS", "THE1 EXCLUSIVE PRICE VND", "NOTE_1"}:
        return "000000", "FFFFFF"
    if column_name == "PAGE DIGITAL":
        return "15536A", "FFFFFF"
    if column_name in {
        "START DATE",
        "END DATE",
        "MULTIBUYCOMBO SELLING PRICE VND",
        "DISCOUNT TO CUSTOMER VND",
    }:
        return "FFF200", "000000"
    if column_name in {
        "SYSTEM GOLD VLL",
        "REGION NORTH SOUTH CENTER ALL",
        "REGION",
        "PAGE",
        "THEMATIC",
        "POSITION",
        "NORMAL SELLING PRICE VND",
        "PROMOTIONAL SELLING PRICE VND",
        "% DISCOUNT TO CUSTOMER",
        "PROMO TYPE",
        "PROMO DETAILS",
        "PROMOTION GIFT",
        "GIFT VALUE",
        "NOTE",
    }:
        return "FF0000", "FFFFFF"
    if column_name == "MDT NOTE":
        return "F4B183", "000000"
    return "B7B7B7", "000000"


def _read_add_workbook(file_bytes: bytes) -> dict[str, pl.DataFrame]:
    sheets = pl.read_excel(BytesIO(file_bytes), sheet_id=0, infer_schema_length=0)
    return sheets if isinstance(sheets, dict) else {"Sheet": sheets}


def build_result_zip(
    results: list[tuple[str, pl.DataFrame]],
    add_files: list[tuple[str, bytes, bool]],
    target_date: date | datetime | str,
) -> tuple[bytes, list[dict[str, object]]]:
    target_date = parse_target_date(target_date)
    archive_buffer = BytesIO()
    package_summary: list[dict[str, object]] = []

    main_by_group: dict[str, list[tuple[str, pl.DataFrame]]] = {
        group: [] for group in BUSINESS_GROUPS
    }
    for sheet_name, dataframe in results:
        for group, subset in split_by_business_group(
            dataframe, "DEPARTMENT", "DIVISION"
        ).items():
            main_by_group[group].append((sheet_name, subset))

    with ZipFile(archive_buffer, "w", compression=ZIP_DEFLATED) as archive:
        main_filename = f"CataPriceCheck_{target_date:%Y%m%d}.xlsx"
        for group, grouped_results in main_by_group.items():
            if not grouped_results:
                continue
            archive.writestr(
                f"{group}/{main_filename}",
                results_to_excel(grouped_results),
            )
            package_summary.append(
                {
                    "FOLDER": group,
                    "FILE": main_filename,
                    "ROWS": sum(frame.height for _, frame in grouped_results),
                }
            )

        for uploaded_name, file_bytes, use_dept in add_files:
            filename = Path(uploaded_name).name
            source_sheets = _read_add_workbook(file_bytes)
            split_sheets: dict[str, list[tuple[str, pl.DataFrame]]] = {
                group: [] for group in BUSINESS_GROUPS
            }

            for sheet_name, dataframe in source_sheets.items():
                primary_column = "DEPT" if use_dept else "DIVISION"
                for group, subset in split_by_business_group(
                    dataframe, primary_column
                ).items():
                    split_sheets[group].append((sheet_name, subset))

            for group, grouped_results in split_sheets.items():
                if not grouped_results:
                    continue
                archive.writestr(
                    f"{group}/{filename}",
                    results_to_excel(grouped_results),
                )
                package_summary.append(
                    {
                        "FOLDER": group,
                        "FILE": filename,
                        "ROWS": sum(frame.height for _, frame in grouped_results),
                    }
                )

    return archive_buffer.getvalue(), package_summary
